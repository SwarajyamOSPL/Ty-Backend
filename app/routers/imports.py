"""Excel import of weekly capper figures.

One workbook, one capper, many weeks: every worksheet is one week. The week is read from a
title cell above the table (e.g. "Baws - October Week 1 (September 28th - October 4th) 2026")
and, if there is none, from the worksheet's name (e.g. "10-6 - 10-12", "Oct 6 - Oct 12",
"106 1011") - see app/week_names.py.

When there is no title text, a date in a cell above the table or in the header row (such as 2026-10-04)
gives the week: the Monday-to-Sunday week that date falls in. Failing that, the worksheet's name is used.

Hidden worksheets are ignored: only the sheets that are visible in Excel are read.

The table can start anywhere: the header row is found by looking for a column called
"username" (or "user", "account", or a header containing the word "user" or "login") and a figure column in the first rows. The figure column is "figure" or, if there is
none, "fig", "total" or "PNL" (in that order of preference); failing that, a column whose header only
CONTAINS one of those words, such as "Weekly Fig" or "Gross PNL" (several with the same word are rejected as
ambiguous). Other columns are ignored.

The file name must contain the capper's name (its color), e.g. "Turquoise.xlsx" or
"TB2 - Yellow.xlsx", and it must be the same capper that was selected.

Rows are matched, by username (ignoring case; numbers like 7261.0 become "7261"), to accounts
already assigned to the selected capper. If several of the capper's accounts have the same
username, a "Website" column in the sheet (when there is one) tells them apart. Every account
that gets a figure has its Site Figure for the week recalculated right away (the sum of all its
assigned cappers' figures, negative values included). Rows whose username is not one of the capper's
accounts are "unregistered rows" (the sheet may contain other partners' accounts that are not in the
database): they are not saved as figures but kept with a note, see GET /cappers/figures/unregistered.
Rows with no username, no figure, or a figure that is not a number are ignored and counted
("ignored_rows"). Worksheets without a username (or account / login)/figure header (notes,
totals...) are ignored and reported. All-or-nothing for real problems: an unreadable week, an
invalid figure or an ambiguous username rejects the whole file. When an account appears in several rows of
a week with different figures, the figures are added up (counted in "summed_rows"); a row repeating a
figure the account already has for that week is ignored (counted in "repeated_ignored").
"""
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Response, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, func, literal, select, union_all
from sqlalchemy.orm import Session

from ..accounts_view import recalc_site_figures
from ..db import get_db
from ..accounts_view import _links
from ..listing import Page, account_conditions, account_ids, contains, page_params, paginate
from ..models import Account, Admin, Figure, Party, PartyAccount, PartyKind, UnregisteredFigure
from ..schemas import ImportedRowOut, ImportResult, SheetResult, UnregisteredFigureOut
from ..security import READ, WRITE
from ..week_names import parse_week
from .parties import CAPPER_COLORS

router = APIRouter(prefix="/cappers/figures/import", tags=["import"])
unregistered_router = APIRouter(prefix="/cappers/figures/unregistered", tags=["import"])
imported_router = APIRouter(prefix="/cappers/figures/imported", tags=["import"])

NOTE_UNREGISTERED = "not registered account"            # the username is not an account in the database
NOTE_NOT_ASSIGNED = "account not assigned to this capper"  # it is an account, but not this capper's

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000          # per worksheet
HEADER_SCAN_ROWS = 30     # the header row must be within the first rows of a worksheet
USERNAME_HEADERS = ("username", "user name", "user", "login", "account")  # in order of preference
FIGURE_HEADERS = ("figure", "fig", "total", "pnl")  # in order of preference


def _norm(v) -> str:
    return str(v).strip().lower() if v is not None else ""


def _text(v) -> str:
    """Cell text for matching: whole numbers stored as floats (7261.0) become '7261'."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if v is not None else ""


def _site(v) -> str:
    """A website for comparing: lowercase, no scheme, no 'www.', no path."""
    s = re.sub(r"^[a-z]+://", "", str(v).strip().lower())
    return re.sub(r"^www\.", "", s).split("/")[0]


def _header(v) -> str:
    """Header text for comparing: lowercase, "_" and runs of spaces become one space."""
    return re.sub(r"[\s_]+", " ", v.strip().lower()) if isinstance(v, str) else ""


def _find_columns(row) -> dict | None:
    """{'username': i, 'figure': j, 'figure_name': text} if this row is a header row.
    The figure column is, in order of preference, one called exactly figure / fig / total / pnl or,
    failing that, one whose header merely CONTAINS one of those words ("Weekly Fig", "Total Figure",
    "Gross PNL", "PNL (USD)"; the word may be plural). Words are tried in the same order. If several
    headers contain the same best word, {'username': i, 'ambiguous': [headers]} is returned."""
    cells = [_header(c) for c in row]
    user_col = next((cells.index(name) for name in USERNAME_HEADERS if name in cells), None)
    for word in ("user", "login"):  # no exact name: the leftmost header that merely contains one of these words
        if user_col is None:
            user_col = next((j for j, c in enumerate(cells) if re.search(rf"\b{word}s?\b", c)), None)
    if user_col is None:
        return None
    website = cells.index("website") if "website" in cells else None

    def found(j):
        return {"username": user_col, "figure": j, "figure_name": str(row[j]).strip(), "website": website}

    for name in FIGURE_HEADERS:
        if name in cells:
            return found(cells.index(name))
    for name in FIGURE_HEADERS:
        hits = [j for j, c in enumerate(cells) if j != user_col and re.search(rf"\b{name}s?\b", c)]
        if len(hits) == 1:
            return found(hits[0])
        if len(hits) > 1:
            return {"username": user_col, "ambiguous": [str(row[j]).strip() for j in hits]}
    return None


def _is_label_row(row) -> bool:
    """A row of only text that has a username header ("Account", "Login"...): the header of a table
    that has no figure column (or none we recognise), which must not be read with the columns of the
    table above it."""
    cells = [_header(c) for c in row if c is not None and str(c).strip()]
    return len(cells) >= 2 and all(isinstance(c, str) for c in row if c is not None) \
        and any(c in USERNAME_HEADERS for c in cells)


def _slug(text: str) -> str:
    """Lowercase, with every run of non-alphanumeric characters turned into one space."""
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _capper_in_filename(filename: str, known_names: list[str]) -> str | None:
    """The capper name found in the file name, or None. Whole words only; if several names
    match (e.g. 'Blue' and 'Light Blue') the longest wins."""
    stem = re.sub(r"\.xlsx$", "", filename.replace("\\", "/").rsplit("/", 1)[-1], flags=re.I)
    padded = f" {_slug(stem)} "
    found = [n for n in known_names if _slug(n) and f" {_slug(n)} " in padded]
    return max(found, key=lambda n: len(_slug(n))) if found else None


def _fail(errors: list[str], status_code=status.HTTP_422_UNPROCESSABLE_ENTITY):
    raise HTTPException(status_code, {"message": "Import rejected, nothing was saved",
                                      "errors": errors[:200]})


def _template_workbook(today: date | None = None) -> Workbook:
    """The capper sheet layout the import reads: one worksheet per week (named by the week, Monday to
    Sunday), a header row with Website / Username / Figure, one row per account. Two empty weeks (this
    week and last week) to start from; add or rename worksheets as needed."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    today = today or date.today()
    monday = today - timedelta(days=today.weekday())
    notes = {
        "Website": "Optional, but recommended: it tells accounts that share a username apart.",
        "Username": "Required. The account's username, one row per account. "
                    "(A header called User, Login or Account also works.)",
        "Figure": "Required. The week's figure for the account; negative numbers are fine. "
                  "(A header containing Fig, Figure, Total or PNL also works.)",
    }
    wb = Workbook()
    wb.remove(wb.active)
    for weeks_back in (0, 1):
        start = monday - timedelta(days=7 * weeks_back)
        end = start + timedelta(days=6)
        ws = wb.create_sheet(f"{start.month}-{start.day} - {end.month}-{end.day}")
        ws.append(list(notes))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(horizontal="center")
            cell.comment = Comment(notes[cell.value], "Capper sheet")
        for col, width in (("A", 28), ("B", 22), ("C", 14)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        check = DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000",
                               allow_blank=True, showErrorMessage=True, errorTitle="Figure",
                               error="The figure must be a number, for example -250.5")
        ws.add_data_validation(check)
        check.add("C2:C5000")
        ws["C2"].number_format = "#,##0.00"
    return wb


@router.get("/template")
def template(_: Admin = Depends(WRITE)):
    """Download the capper sheet template (.xlsx): worksheets named by week, columns Website,
    Username, Figure. Fill in one row per account, add a worksheet for each week (name it like
    "10-6 - 10-12"), and name the file after the capper (e.g. "Pink.xlsx") before uploading."""
    buf = io.BytesIO()
    _template_workbook().save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="capper_figures_template.xlsx"'})


@router.post("", response_model=ImportResult)
async def import_figures(capper_id: int = Form(...), file: UploadFile = File(...),
                         db: Session = Depends(get_db), user: Admin = Depends(WRITE)):
    """Each worksheet is one week. The dates come from a title cell above the table, or from the
    worksheet's name. When the dates have no year it is the latest year in which the week does
    not start more than 7 days after today."""
    capper = db.get(Party, capper_id)
    if not capper or capper.kind != PartyKind.capper:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Capper not found")
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Upload an .xlsx file")

    known = list({*CAPPER_COLORS, *db.scalars(select(Party.name).where(Party.kind == PartyKind.capper))})
    in_name = _capper_in_filename(file.filename, known)
    if not in_name:
        _fail([f"No capper name found in the file name '{file.filename}'. Name the file after "
               f"the selected capper, e.g. '{capper.name}.xlsx'"])
    if _slug(in_name) != _slug(capper.name):
        _fail([f"File name is for capper '{in_name}' but the selected capper is '{capper.name}'"])

    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "File larger than 5 MB")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]  # hidden worksheets are never read
        hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    except Exception:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Not a valid .xlsx file")

    # accounts assigned to this capper, indexed by username
    assigned = db.scalars(select(Account).join(PartyAccount, PartyAccount.account_id == Account.id)
                          .where(PartyAccount.party_id == capper.id)).all()
    by_user: dict[str, list[Account]] = {}
    for a in assigned:
        if _norm(a.username):
            by_user.setdefault(_norm(a.username), []).append(a)

    all_usernames = {_norm(u) for u in db.scalars(select(Account.username))}  # every account in the database
    errors: list[str] = []
    ignored: list[str] = [f"{t} (hidden)" for t in hidden]
    weeks: dict[tuple, dict] = {}  # (start, end) -> data collected from that worksheet
    for ws in sheets:
        # find the header row (within the first rows) and remember the text above it
        try:
            rows = ws.iter_rows(values_only=True)
            above: list[str] = []
            dates: list[date] = []  # date cells in the rows above the table and in the header row
            cols = None
            header_row_no = 0
            for header_row_no, row in enumerate(rows, start=1):
                cols = _find_columns(row)
                dates.extend(c.date() if isinstance(c, datetime) else c for c in row if isinstance(c, (datetime, date)))
                if cols or header_row_no >= HEADER_SCAN_ROWS:
                    break
                above.extend(c.strip() for c in row if isinstance(c, str) and c.strip())
        except Exception:
            errors.append(f"[{ws.title}] can't be read")
            continue
        if not cols:
            ignored.append(ws.title)  # not a data sheet (notes, totals, ...)
            continue
        if cols.get("ambiguous"):
            errors.append(f"[{ws.title}] has several figure columns ({', '.join(cols['ambiguous'])}) and "
                          f"I can't tell which one is the figure. Keep one of them, or name it exactly 'Figure'")
            continue

        # the week: a title cell above the table first, then the worksheet's name
        week, week_from = None, ""
        for text in above:
            try:
                week, week_from = parse_week(text), "sheet title"
                break
            except ValueError:
                continue
        if week is None and dates:
            # a date cell (e.g. 2026-10-04 in the header row): the Monday-Sunday week it falls in
            monday = dates[0] - timedelta(days=dates[0].weekday())
            week, week_from = (monday, monday + timedelta(days=6)), "date cell"
        if week is None:
            try:
                week, week_from = parse_week(ws.title), "sheet name"
            except ValueError as e:
                errors.append(f"{e}. Name each sheet by its week, e.g. '10-6 - 10-12', "
                              f"'Oct 6 - Oct 12' or 'April w2', or put the week in a title cell above the table")
                continue
        if week in weeks:
            errors.append(f"[{ws.title}] is the same week ({week[0]} to {week[1]}) as "
                          f"[{weeks[week]['sheet']}]")
            continue

        entries: dict[int, Decimal] = {}  # account_id -> figure
        ignored_rows = 0
        repeated_rows = 0
        summed_rows = 0
        seen: dict[int, set[Decimal]] = {}  # account_id -> the distinct figures already read for it
        unregistered: dict[tuple, dict] = {}  # rows of accounts that are not the capper's, kept with a note
        n_rows = 0
        for n, row in enumerate(rows, start=header_row_no + 1):
            # A sheet may stack several tables with different columns (e.g. "Book | Account | Figure"
            # and, further down, "Author | Book | Account | Password | Figure"): every header row
            # switches the columns used for the rows below it. A table without a figure column is
            # skipped until the next header.
            new = _find_columns(row)
            if new or _is_label_row(row):
                if new and new.get("ambiguous"):
                    errors.append(f"[{ws.title}] row {n}: several figure columns ({', '.join(new['ambiguous'])}) "
                                  f"and I can't tell which one is the figure")
                cols = new if new and not new.get("ambiguous") else None
                continue
            if cols is None:
                continue
            u_raw = row[cols["username"]] if cols["username"] < len(row) else None
            raw = row[cols["figure"]] if cols["figure"] < len(row) else None
            site_raw = row[cols["website"]] if cols["website"] is not None and cols["website"] < len(row) else None
            if (u_raw is None or str(u_raw).strip() == "") and (raw is None or str(raw).strip() == ""):
                continue  # blank line
            n_rows += 1
            if n_rows > MAX_ROWS:
                errors.append(f"[{ws.title}] more than {MAX_ROWS} rows")
                break
            user_ = _norm(_text(u_raw))
            matches = by_user.get(user_, []) if user_ else []
            blank_figure = raw is None or str(raw).strip() == ""
            if not matches or blank_figure:
                # Not one of this capper's accounts (e.g. another partner's account that is not in the
                # database) -> an "unregistered row", kept with a note when it has a figure; or no
                # usable figure / username -> ignored. Either way the row is reported.
                fig = None
                if not blank_figure and not isinstance(raw, bool):
                    try:
                        fig = Decimal(str(raw).strip().replace(",", ""))
                        if not fig.is_finite() or abs(fig) >= Decimal("1e12"):
                            fig = None
                    except InvalidOperation:
                        fig = None
                if not user_ or fig is None or matches:
                    ignored_rows += 1  # no username, no figure, or a figure that is not a number
                    continue
                site = str(site_raw).strip()[:255] if site_raw is not None and str(site_raw).strip() else None
                key = (user_, _site(site) if site else "", fig)
                if key in unregistered:
                    repeated_rows += 1  # the same unregistered row again
                    continue
                unregistered[key] = {"username": _text(u_raw)[:255], "website": site, "figure": fig,
                                     "notes": NOTE_NOT_ASSIGNED if user_ in all_usernames else NOTE_UNREGISTERED}
                continue
            try:
                if isinstance(raw, bool):
                    raise InvalidOperation
                figure = Decimal(str(raw).strip().replace(",", ""))
                if not figure.is_finite() or abs(figure) >= Decimal("1e12"):
                    raise InvalidOperation
            except InvalidOperation:
                errors.append(f"[{ws.title}] row {n}: figure '{raw}' is not a valid number")
                continue
            if len(matches) > 1 and site_raw is not None and str(site_raw).strip():
                narrowed = [a for a in matches if _site(site_raw) in {_site(w) for w in a.websites}]
                if len(narrowed) == 1:  # same username on several sites: the sheet's website decides
                    matches = narrowed
            if len(matches) > 1:
                errors.append(f"[{ws.title}] row {n}: {len(matches)} accounts of capper "
                              f"{capper.name} have username '{user_}'"
                              + ("" if site_raw is not None else " (add a Website column to tell them apart)"))
                continue
            acct = matches[0]
            if acct.id in entries:
                if figure in seen[acct.id]:
                    repeated_rows += 1  # the same account with the same figure again: just ignore the repeat
                    continue
                entries[acct.id] += figure  # another figure for the same account: they add up
                summed_rows += 1
            else:
                entries[acct.id] = figure
            seen.setdefault(acct.id, set()).add(figure)
        weeks[week] = {"sheet": ws.title, "entries": entries, "ignored_rows": ignored_rows,
                       "repeated_rows": repeated_rows, "summed_rows": summed_rows,
                       "unregistered": list(unregistered.values()), "week_from": week_from,
                       "figure_column": cols["figure_name"], "header_row": header_row_no}

    if not weeks and not errors:
        _fail(["No worksheet has a header row with a 'username' (or 'user' / 'account' / a header containing 'user' or 'login') column and a figure column "
               "(a header containing 'figure', 'fig', 'total' or 'PNL', e.g. 'Weekly Fig' or 'Gross PNL')"
               + (f" (ignored sheets: {', '.join(ignored)})" if ignored else "")])
    if errors:
        _fail(errors)
    if not any(w["entries"] for w in weeks.values()):
        total_skipped = sum(w["ignored_rows"] + len(w["unregistered"]) for w in weeks.values())
        _fail([f"No row matched an account assigned to capper {capper.name} "
               f"({total_skipped} row(s) not saved). Check that this is the right capper's file "
               f"and that the accounts are assigned to the capper"])

    # Re-importing the same week overwrites instead of duplicating.
    results: list[SheetResult] = []
    for (start, end), w in weeks.items():
        entries = w["entries"]
        existing = {f.account_id: f for f in db.scalars(select(Figure).where(
            Figure.party_id == capper.id, Figure.start_date == start,
            Figure.end_date == end, Figure.account_id.in_(entries)))} if entries else {}
        created = updated = 0
        for account_id, figure in entries.items():
            fig = existing.get(account_id)
            if fig:
                fig.figure, fig.entered_by = figure, user.id
                updated += 1
            else:
                db.add(Figure(party_id=capper.id, account_id=account_id, figure=figure,
                              start_date=start, end_date=end, entered_by=user.id))
                created += 1
        # rows that could not be saved are kept with a note; the same capper and week replace the old ones
        db.execute(delete(UnregisteredFigure).where(
            UnregisteredFigure.capper_id == capper.id, UnregisteredFigure.start_date == start,
            UnregisteredFigure.end_date == end))
        for u in w["unregistered"]:
            db.add(UnregisteredFigure(capper_id=capper.id, start_date=start, end_date=end,
                                      imported_by=user.id, **u))
        db.flush()
        site_figures = recalc_site_figures(db, list(entries), (start, end), user.id)
        results.append(SheetResult(sheet_name=w["sheet"], start_date=start, end_date=end,
                                   week_from=w["week_from"], header_row=w["header_row"],
                                   figure_column=w["figure_column"],
                                   created=created, updated=updated, site_figures_updated=site_figures,
                                   unregistered_rows=len(w["unregistered"]), ignored_rows=w["ignored_rows"],
                                   repeated_ignored=w["repeated_rows"], summed_rows=w["summed_rows"]))
    db.commit()
    return ImportResult(capper_id=capper.id, capper_name=capper.name,
                        created=sum(r.created for r in results), updated=sum(r.updated for r in results),
                        unregistered_rows=sum(r.unregistered_rows for r in results),
                        ignored_rows=sum(r.ignored_rows for r in results),
                        site_figures_updated=sum(r.site_figures_updated for r in results),
                        sheets=results, ignored_sheets=ignored)


@unregistered_router.get("", response_model=list[UnregisteredFigureOut], dependencies=[Depends(READ)])
def list_unregistered(response: Response, capper_id: int | None = None, start_date: date | None = None,
                      end_date: date | None = None, username: str | None = None, website: str | None = None,
                      notes: str | None = None, pg: Page = Depends(page_params), db: Session = Depends(get_db)):
    """Rows of imported capper sheets that were not saved as figures, with the note why:
    "not registered account" (no such account in the database) or "account not assigned to this
    capper". Newest week first, one page at a time. `username`, `website` and `notes` are
    case-insensitive "contains" filters."""
    q = (select(UnregisteredFigure).order_by(UnregisteredFigure.start_date.desc(), UnregisteredFigure.id))
    if capper_id:
        q = q.where(UnregisteredFigure.capper_id == capper_id)
    if start_date:
        q = q.where(UnregisteredFigure.start_date >= start_date)
    if end_date:
        q = q.where(UnregisteredFigure.end_date <= end_date)
    if username and username.strip():
        q = q.where(contains(UnregisteredFigure.username, username))
    if website and website.strip():
        q = q.where(contains(UnregisteredFigure.website, website))
    if notes and notes.strip():
        q = q.where(contains(UnregisteredFigure.notes, notes))
    rows = paginate(db, q, response, pg)
    return [UnregisteredFigureOut(id=r.id, capper_id=r.capper_id, capper_name=r.capper.name,
                                  username=r.username, website=r.website, figure=r.figure,
                                  start_date=r.start_date, end_date=r.end_date, notes=r.notes,
                                  created_at=r.created_at) for r in rows]


@imported_router.get("", response_model=list[ImportedRowOut], dependencies=[Depends(READ)])
def list_imported(response: Response, capper_id: int | None = None, start_date: date | None = None,
                  end_date: date | None = None, username: str | None = None, website: str | None = None,
                  status_: str = Query("all", alias="status", pattern="^(all|registered|unregistered)$",
                                       description="all, registered (saved figures) or unregistered"),
                  pg: Page = Depends(page_params), db: Session = Depends(get_db)):
    """The data of imported capper sheets in one list: the figures saved for registered accounts and
    the rows kept because the account is not registered (or not the capper's), each with its week and
    `notes` (null when the figure was saved). Newest week first, one page at a time (see
    app/listing.py). Filters: `capper_id`, `start_date`, `end_date`, `status`, and case-insensitive
    "contains" filters `username` and `website`."""
    parts = []
    if status_ in ("all", "registered"):
        q = (select(Figure.start_date.label("sd"), literal(0).label("k"), Figure.id.label("rid"),
                    Account.username.label("uname"))
             .select_from(Figure).join(Party, Party.id == Figure.party_id)
             .join(Account, Account.id == Figure.account_id).where(Party.kind == PartyKind.capper))
        if capper_id:
            q = q.where(Figure.party_id == capper_id)
        if start_date:
            q = q.where(Figure.start_date >= start_date)
        if end_date:
            q = q.where(Figure.end_date <= end_date)
        conds = account_conditions(username=username, website=website)
        if conds:
            q = q.where(Figure.account_id.in_(account_ids(conds)))
        parts.append(q)
    if status_ in ("all", "unregistered"):
        q = select(UnregisteredFigure.start_date.label("sd"), literal(1).label("k"),
                   UnregisteredFigure.id.label("rid"), UnregisteredFigure.username.label("uname"))
        if capper_id:
            q = q.where(UnregisteredFigure.capper_id == capper_id)
        if start_date:
            q = q.where(UnregisteredFigure.start_date >= start_date)
        if end_date:
            q = q.where(UnregisteredFigure.end_date <= end_date)
        if username and username.strip():
            q = q.where(contains(UnregisteredFigure.username, username))
        if website and website.strip():
            q = q.where(contains(UnregisteredFigure.website, website))
        parts.append(q)
    u = union_all(*parts).subquery()
    total = db.scalar(select(func.count()).select_from(u)) or 0
    response.headers["X-Total-Count"] = str(total)
    response.headers["X-Page"] = str(pg.page)
    response.headers["X-Page-Size"] = str(pg.size)
    response.headers["X-Total-Pages"] = str(-(-total // pg.size))
    page = db.execute(select(u.c.k, u.c.rid).order_by(u.c.sd.desc(), u.c.k, func.lower(u.c.uname), u.c.rid)
                      .offset((pg.page - 1) * pg.size).limit(pg.size)).all()
    reg_ids = [rid for k, rid in page if k == 0]
    unr_ids = [rid for k, rid in page if k == 1]
    figs = {f.id: f for f in db.scalars(select(Figure).where(Figure.id.in_(reg_ids)))} if reg_ids else {}
    unrs = {r.id: r for r in db.scalars(select(UnregisteredFigure).where(UnregisteredFigure.id.in_(unr_ids)))} if unr_ids else {}
    _, partners = _links(db, sorted({f.account_id for f in figs.values()}))
    out = []
    for k, rid in page:
        if k == 0:
            f = figs[rid]
            out.append(ImportedRowOut(
                source="registered", id=f.id, capper_id=f.party_id, capper_name=f.party.name,
                account_id=f.account_id, username=f.account.username, websites=f.account.websites,
                partners=partners.get(f.account_id, []), figure=f.figure, start_date=f.start_date,
                end_date=f.end_date, notes=None, created_at=f.created_at))
        else:
            r = unrs[rid]
            out.append(ImportedRowOut(
                source="unregistered", id=r.id, capper_id=r.capper_id, capper_name=r.capper.name,
                account_id=None, username=r.username, websites=[r.website] if r.website else [],
                partners=[], figure=r.figure, start_date=r.start_date, end_date=r.end_date,
                notes=r.notes, created_at=r.created_at))
    return out

