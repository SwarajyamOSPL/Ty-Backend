"""Import: headers anywhere, column aliases, numeric usernames, week from a title cell."""
import io
import os

os.environ["DATABASE_URL"] = "sqlite:///./test.db"

from datetime import date  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.security import seed_admin  # noqa: E402
from app.week_names import parse_week  # noqa: E402


def make_client():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_admin(db)
    c = TestClient(app)
    tok = c.post("/auth/login", data={"username": "admin", "password": "admin"}).json()
    return c, {"Authorization": f"Bearer {tok['access_token']}"}


def sheet_file(sheets):
    """sheets: {sheet name: list of raw rows}; rows are written as given, starting at row 1."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_title_with_week_number_and_year():
    t = "Baws - October Week 1 (September 28th - October 4th) 2026"
    assert parse_week(t, date(2026, 10, 13)) == (date(2026, 9, 28), date(2026, 10, 4))
    # the same text from a later year-less read still takes the year written in the title
    assert parse_week(t, date(2030, 1, 1)) == (date(2026, 9, 28), date(2026, 10, 4))


def test_real_layout(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    accts = {}
    for site, user in (("trylnk.com", "lnk98213"), ("empire.ag", "Ccm166"), ("sports411.ag", "8548"),
                       ("dup.com", "twin"), ("dup2.com", "TWIN")):
        a = c.post("/accounts", json={"websites": [site], "username": user}, headers=h).json()
        accts[user] = a
        c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": a["id"]}, headers=h)

    def upload(content, name="TB2 - Yellow.xlsx"):
        return c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
                      files={"file": (name, content)}, headers=h)

    def sheet_info(r):
        return r.json()["sheets"][0]

    title = "Baws - October Week 1 (September 28th - October 4th) 2026"
    body = [("legacy365.net", 7261.0, 0.0, "x"), ("sports411.ag", 8548.0, 150.5, "x"),
            ("empire.ag", "CCM166", -20.25, "x"), ("trylnk.com", "LNK98213", None, "x")]
    # header on row 7 in column C..F, a title above it, week number in the title, sheet named "104"
    rows = [[], [], [], [], [], [None, None, title], [None, None, "Website", "Username", "TOTAL", 32360.83]]
    rows += [[None, None, *r] for r in body]
    r = upload(sheet_file({"104": rows}))
    assert r.status_code == 200, r.text
    s = sheet_info(r)
    assert (s["week_from"], s["header_row"], s["figure_column"]) == ("sheet title", 7, "TOTAL")
    assert (s["start_date"], s["end_date"]) == ("2026-09-28", "2026-10-04")  # not the sheet name "104" (Oct 4 + 6)
    assert (s["created"], s["unregistered_rows"], s["ignored_rows"]) == (2, 1, 1)  # 8548.0 matches; 7261 is unknown; a blank figure is ignored
    saved = {f["username"]: f["figure"] for f in c.get("/cappers/figures/all", headers=h).json()}
    assert saved == {"8548": "150.50", "Ccm166": "-20.25"}

    # the title wins over the sheet name; with no title the sheet name is used
    plain = [["Username", "Total"], ["lnk98213", 9]]
    r = upload(sheet_file({"10-13 - 10-19": [[title], *plain]}))
    assert sheet_info(r)["start_date"] == "2026-09-28" and sheet_info(r)["week_from"] == "sheet title"
    r = upload(sheet_file({"10-13 - 10-19": plain}))
    assert (sheet_info(r)["start_date"], sheet_info(r)["week_from"]) == ("2026-10-13", "sheet name")

    # figure column names: figure > fig > total > pnl; any column position; any case
    wk = "10-13 - 10-19"
    for header, want in ((["x", "PNL", "Username"], "PNL"), (["Username", "TOTAL", "Fig"], "Fig"),
                         (["Username", "total", "PNL", "FIGURE"], "FIGURE"), (["USERNAME", "pnl"], "pnl"),
                         ([" User Name ", "  Fig  "], "Fig")):
        u_i = [i for i, x in enumerate(header) if x.strip().lower() in ("username", "user name")][0]
        want_i = [i for i, x in enumerate(header) if x.strip() == want][0]
        row = [None] * len(header)
        row[u_i] = "lnk98213"
        row[want_i] = 5
        for i, x in enumerate(header):  # other candidate columns get a different number
            if i not in (u_i, want_i) and x.strip().lower() in ("fig", "total", "pnl", "figure"):
                row[i] = 999
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and sheet_info(r)["figure_column"] == want, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "5.00", header

    # the username column may also be called exactly "Account" (any case, spaces around it)
    for header in (["Account", "PNL"], ["ACCOUNT", "Total"], ["Website", " account ", "Figure"]):
        u_i = next(i for i, x in enumerate(header) if x.strip().lower() == "account")
        f_i = next(i for i, x in enumerate(header) if x.strip().lower() in ("pnl", "total", "fig", "figure"))
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 7
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and r.json()["created"] + r.json()["updated"] == 1, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "7.00", header
    # "Account Name" is not a username column any more: strictly "Account"
    for header in (["Account Name", "Total"], ["account_name", "fig"], ["AccountName", "Total"], ["Account No", "Total"]):
        assert upload(sheet_file({wk: [header, ["lnk98213", 1]]})).status_code == 422, header
    # "Login" also works as the username column (exact header only)
    for header in (["Login", "PNL"], ["LOGIN", "Total"], ["Website", " login ", "Figure"]):
        u_i = next(i for i, x in enumerate(header) if x.strip().lower() == "login")
        f_i = next(i for i, x in enumerate(header) if x.strip().lower() in ("pnl", "total", "figure"))
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 4
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and r.json()["created"] + r.json()["updated"] == 1, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "4.00", header
    # a header that merely CONTAINS the word login is a username column too (whole word, any case)
    for header in (["Login ID", "Total"], ["Player Login Name", "Fig"], ["Website", "login_id", "Figure"],
                   ["LOGIN/USERNAME", "PNL"]):
        u_i = next(i for i, x in enumerate(header) if "login" in x.lower())
        f_i = next(i for i, x in enumerate(header) if x.strip().lower() in ("pnl", "total", "fig", "figure"))
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 8
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and r.json()["created"] + r.json()["updated"] == 1, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "8.00", header
    for header in (["Relogin", "Total"], ["Loginx", "Total"], ["Blogin", "Total"]):   # pieces of a word do not count
        assert upload(sheet_file({wk: [header, ["lnk98213", 1]]})).status_code == 422, header
    # "User" is a username column, and so is any header that contains the word user (also plural);
    # a header containing "user" is preferred to one containing "login"
    for header in (["User", "Total"], ["USER", "PNL"], ["Website", "User ID", "Figure"], ["Player User", "Fig"],
                   ["Users", "Total"], ["Login Date", "User ID", "Total"]):
        u_i = next(i for i, x in enumerate(header) if "user" in x.lower())
        f_i = next(i for i, x in enumerate(header) if x.strip().lower() in ("pnl", "total", "fig", "figure"))
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 3
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and r.json()["created"] + r.json()["updated"] == 1, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "3.00", header
    for header in (["Superuser", "Total"], ["Userx", "Total"]):   # pieces of a word do not count
        assert upload(sheet_file({wk: [header, ["lnk98213", 1]]})).status_code == 422, header
    # an exact name beats a header that contains login
    r = upload(sheet_file({wk: [["Login Date", "Username", "Total"], ["2026-10-01", "lnk98213", 9]]}))
    assert r.status_code == 200 and c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "9.00"
    r = upload(sheet_file({wk: [["Login", "Username", "Total"], ["no_such_account", "lnk98213", 6]]}))  # "username" wins
    assert r.status_code == 200 and c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "6.00"
    # "Account" is a username column too, but "username" / "login" win over it
    assert upload(sheet_file({wk: [["Account", "Total"], ["lnk98213", 1]]})).status_code == 200
    r = upload(sheet_file({wk: [["Account", "Username", "Total"], ["no_such_account", "lnk98213", 12]]}))
    assert r.status_code == 200 and c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "12.00"
    # if both exist, "username" wins and the other column is ignored
    r = upload(sheet_file({wk: [["Account Name", "Username", "Total"], ["no_such_account", "lnk98213", 11]]}))
    assert r.status_code == 200 and c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "11.00"

    # a header that merely CONTAINS the word PNL counts as the figure column
    for header, want in ((["Username", "Gross PNL"], "Gross PNL"), (["Gross_PNL", "Account"], "Gross_PNL"),
                         (["Username", "Net PNL (USD)"], "Net PNL (USD)"), (["Username", "weekly pnl", "Website"], "weekly pnl"),
                         (["Username", "Gross PNL", "PNL"], "PNL"),          # an exact name beats a "contains PNL" header
                         (["Username", "Gross PNL", "Total"], "Total")):
        u_i = next(i for i, x in enumerate(header) if x.lower() in ("username", "account"))
        f_i = header.index(want)
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 9
        for i, x in enumerate(header):  # the other candidate columns hold a different number
            if i not in (u_i, f_i) and "pnl" in x.lower():
                row[i] = 999
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and sheet_info(r)["figure_column"] == want, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "9.00", header
    # any header that contains one of the words counts, in any case/position, also plural: figure, fig, total, pnl
    for header, want in ((["Username", "Weekly Fig"], "Weekly Fig"), (["Username", "Total Figure"], "Total Figure"),
                         (["Username", "Site Fig"], "Site Fig"), (["Username", "Figs"], "Figs"),
                         (["Username", "Week Total", "Notes"], "Week Total"),
                         (["Username", "Weekly PNL", "Fig Total"], "Fig Total")):   # figure > fig > total > pnl
        u_i, f_i = header.index("Username"), header.index(want)
        row = [None] * len(header)
        row[u_i], row[f_i] = "lnk98213", 7
        r = upload(sheet_file({wk: [header, row]}))
        assert r.status_code == 200 and sheet_info(r)["figure_column"] == want, (header, r.text)
        assert c.get("/cappers/figures/all", headers=h).json()[0]["figure"] == "7.00", header
    # but only whole words: pieces of a word do not count
    for header in (["Username", "Pnlx"], ["Username", "Unpnl"], ["Username", "Subtotal"], ["Username", "Configuration"]):
        assert upload(sheet_file({wk: [header, ["lnk98213", 1]]})).status_code == 422, header
    # several columns with the same best word and nothing exact: refuse to guess
    r = upload(sheet_file({wk: [["Username", "Gross PNL", "Net PNL"], ["lnk98213", 1, 2]]}))
    assert r.status_code == 422 and "Gross PNL" in r.json()["detail"]["errors"][0] and "Net PNL" in r.json()["detail"]["errors"][0]
    r = upload(sheet_file({wk: [["Username", "Weekly Fig", "Old Fig"], ["lnk98213", 1, 2]]}))
    assert r.status_code == 422 and "Weekly Fig" in r.json()["detail"]["errors"][0] and "Old Fig" in r.json()["detail"]["errors"][0]

    # no figure-like column, or no username column: the sheet is not a data sheet
    assert upload(sheet_file({wk: [["Username", "Amount"], ["a", 1]]})).status_code == 422
    assert upload(sheet_file({wk: [["Website", "Total"], ["a", 1]]})).status_code == 422
    # header too far down is not found
    assert upload(sheet_file({wk: [[]] * 40 + [["Username", "Total"], ["lnk98213", 1]]})).status_code == 422

    # a username shared by two of the capper's accounts is ambiguous, in any case
    r = upload(sheet_file({wk: [["Username", "Total"], ["Twin", 1]]}))
    assert r.status_code == 422 and "twin" in r.json()["detail"]["errors"][0]


def test_hidden_worksheets_are_ignored(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    a = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": a["id"]}, headers=h)

    # one visible week and two hidden ones; the hidden ones have the same week as each other (would be refused)
    wb = Workbook()
    wb.remove(wb.active)
    for name, state, fig in (("10-13 - 10-19", "visible", 5), ("10-6 - 10-12", "hidden", 99), ("10-6 - 10-12 (2)", "hidden", 77)):
        ws = wb.create_sheet(name)
        ws.sheet_state = state
        ws.append(["Username", "Total"])
        ws.append(["lnk98213", fig])
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
               files={"file": ("Yellow.xlsx", buf.getvalue())}, headers=h)
    assert r.status_code == 200, r.text
    assert [s["sheet_name"] for s in r.json()["sheets"]] == ["10-13 - 10-19"]
    saved = c.get("/cappers/figures/all", headers=h).json()
    assert [(f["figure"], f["start_date"]) for f in saved] == [("5.00", "2026-10-13")]


def test_week_from_a_date_cell(monkeypatch):
    import app.routers.imports as imports_module
    from datetime import datetime

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    a = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": a["id"]}, headers=h)

    def upload(sheets):
        return c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
                      files={"file": ("Yellow.xlsx", sheet_file(sheets))}, headers=h)

    # the date sits in the header row (like "104": 2026-10-04, a Sunday) -> the Monday-Sunday week it is in
    r = upload({"104": [[datetime(2026, 10, 4), "Website", "Account", "Password", "Net Figure"],
                        [None, "trylnk.com", "lnk98213", "x", -490]]})
    assert r.status_code == 200, r.text
    s = r.json()["sheets"][0]
    assert (s["start_date"], s["end_date"], s["week_from"]) == ("2026-09-28", "2026-10-04", "date cell")
    # a Monday date is the first day of its week
    r = upload({"Sheet2": [[datetime(2026, 3, 2), "Account", "Net Figure"], [None, "lnk98213", 5]]})
    assert (r.json()["sheets"][0]["start_date"], r.json()["sheets"][0]["end_date"]) == ("2026-03-02", "2026-03-08")
    # title text still wins over a date cell
    r = upload({"x": [["Baws - October Week 1 (September 28th - October 4th) 2026"],
                      [datetime(2026, 3, 2), "Account", "Net Figure"], [None, "lnk98213", 6]]})
    assert r.json()["sheets"][0]["week_from"] == "sheet title"


def test_unregistered_rows_are_stored_with_a_note(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    blue = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    mine = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213"}, headers=h).json()
    other = c.post("/accounts", json={"websites": ["empire.ag"], "username": "Ccm166"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": mine["id"]}, headers=h)
    c.post(f"/cappers/{blue['id']}/accounts", json={"account_id": other["id"]}, headers=h)

    rows = [["Website", "Account", "Total"],
            ["trylnk.com", "lnk98213", 10],          # saved
            ["empire.ag", "CCM166", 20],             # exists, but is Blue's account
            ["nowhere.com", "ghost1", -5.5],         # no such account at all
            ["nowhere.com", "ghost1", -5.5],         # the same row again: kept once
            ["nowhere.com", "ghost2", None],         # no figure: nothing to keep
            ["nowhere.com", "ghost3", "abc"]]        # not a number: skipped as before, nothing kept
    wk = "10-13 - 10-19"

    def upload(r):
        return c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
                      files={"file": ("Yellow.xlsx", sheet_file({wk: r}))}, headers=h)

    r = upload(rows)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 1 and r.json()["unregistered_rows"] == 2 and r.json()["ignored_rows"] == 2
    s0 = r.json()["sheets"][0]
    assert (s0["unregistered_rows"], s0["ignored_rows"], s0["repeated_ignored"]) == (2, 2, 1)
    assert "skipped" not in r.json() and "skipped" not in s0 and "skipped_usernames" not in s0 and "unregistered_usernames" not in s0

    got = c.get("/cappers/figures/unregistered", headers=h)
    assert got.status_code == 200 and got.headers["x-total-count"] == "2"
    by = {x["username"].lower(): x for x in got.json()}
    assert by["ghost1"]["notes"] == "not registered account" and by["ghost1"]["figure"] == "-5.50"
    assert by["ghost1"]["website"] == "nowhere.com" and by["ghost1"]["capper_name"] == "Yellow"
    assert (by["ghost1"]["start_date"], by["ghost1"]["end_date"]) == ("2026-10-13", "2026-10-19")
    assert by["ccm166"]["notes"] == "account not assigned to this capper" and by["ccm166"]["figure"] == "20.00"
    # filters
    assert [x["username"] for x in c.get("/cappers/figures/unregistered", params={"notes": "not registered"}, headers=h).json()] == ["ghost1"]
    assert len(c.get("/cappers/figures/unregistered", params={"username": "ccm"}, headers=h).json()) == 1
    assert c.get("/cappers/figures/unregistered", params={"capper_id": blue["id"]}, headers=h).json() == []

    # importing the same capper and week again replaces the stored rows instead of adding more
    r = upload(rows)
    assert r.status_code == 200 and c.get("/cappers/figures/unregistered", headers=h).headers["x-total-count"] == "2"
    r = upload([["Website", "Account", "Total"], ["trylnk.com", "lnk98213", 11], ["x.com", "ghost9", 3]])
    assert [x["username"] for x in c.get("/cappers/figures/unregistered", headers=h).json()] == ["ghost9"]
    # no login -> no access
    assert c.get("/cappers/figures/unregistered").status_code in (401, 403)


def test_imported_list_has_registered_and_unregistered_rows(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    blue = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    partner = c.post("/partners", json={"name": "Baws"}, headers=h).json()
    mine = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213",
                                     "partner_ids": [partner["id"]]}, headers=h).json()
    other = c.post("/accounts", json={"websites": ["empire.ag"], "username": "Ccm166"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": mine["id"]}, headers=h)
    c.post(f"/cappers/{blue['id']}/accounts", json={"account_id": other["id"]}, headers=h)

    def upload(wk, rows):
        return c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
                      files={"file": ("Yellow.xlsx", sheet_file({wk: rows}))}, headers=h)

    head = ["Website", "Account", "Total"]
    assert upload("10-13 - 10-19", [head, ["trylnk.com", "lnk98213", 10], ["empire.ag", "CCM166", 20],
                                    ["x.com", "ghost1", -5.5]]).status_code == 200
    assert upload("10-6 - 10-12", [head, ["trylnk.com", "lnk98213", 7], ["x.com", "ghost2", 1]]).status_code == 200

    r = c.get("/cappers/figures/imported", params={"capper_id": yellow["id"]}, headers=h)
    assert r.status_code == 200 and r.headers["x-total-count"] == "5"
    rows = r.json()
    # newest week first, then registered before unregistered, then by username
    assert [(x["start_date"], x["source"], x["username"].lower()) for x in rows] == [
        ("2026-10-13", "registered", "lnk98213"), ("2026-10-13", "unregistered", "ccm166"),
        ("2026-10-13", "unregistered", "ghost1"), ("2026-10-06", "registered", "lnk98213"),
        ("2026-10-06", "unregistered", "ghost2")]
    reg, unr = rows[0], rows[2]
    assert reg["figure"] == "10.00" and reg["notes"] is None and reg["account_id"] == mine["id"]
    assert reg["websites"] == ["trylnk.com"] and [p["name"] for p in reg["partners"]] == ["Baws"]
    assert reg["capper_name"] == "Yellow" and reg["end_date"] == "2026-10-19"
    assert unr["notes"] == "not registered account" and unr["account_id"] is None and unr["figure"] == "-5.50"
    assert unr["websites"] == ["x.com"] and unr["partners"] == []
    assert rows[1]["notes"] == "account not assigned to this capper"
    # filters and paging
    q = lambda **p: c.get("/cappers/figures/imported", params={"capper_id": yellow["id"], **p}, headers=h)
    assert [x["source"] for x in q(status="registered").json()] == ["registered", "registered"]
    assert {x["source"] for x in q(status="unregistered").json()} == {"unregistered"} and q(status="unregistered").headers["x-total-count"] == "3"
    assert [x["username"] for x in q(username="ghost").json()] == ["ghost1", "ghost2"]
    assert len(q(start_date="2026-10-13").json()) == 3
    p2 = q(page=2, page_size=2)
    assert p2.headers["x-total-pages"] == "3" and [x["username"].lower() for x in p2.json()] == ["ghost1", "lnk98213"]
    assert c.get("/cappers/figures/imported", params={"capper_id": blue["id"]}, headers=h).json() == []
    assert q(status="bogus").status_code == 422
    assert c.get("/cappers/figures/imported").status_code in (401, 403)


def test_sheet_named_like_april_w2(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    a = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": a["id"]}, headers=h)
    r = c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
               files={"file": ("Yellow.xlsx", sheet_file({"April w2": [["Account", "Net Figure"], ["lnk98213", 5]],
                                                          "April w3": [["Account", "Net Figure"], ["lnk98213", 6]]}))},
               headers=h)
    assert r.status_code == 200, r.text
    got = {s["sheet_name"]: (s["start_date"], s["end_date"], s["week_from"]) for s in r.json()["sheets"]}
    assert got == {"April w2": ("2026-04-06", "2026-04-12", "sheet name"),
                   "April w3": ("2026-04-13", "2026-04-19", "sheet name")}


def test_stacked_tables_in_one_sheet():
    """Several tables with different columns under each other: each header row switches the columns."""
    c, h = make_client()
    teal = c.post("/cappers", json={"name": "Teal"}, headers=h).json()
    for user in ("a1", "b2", "c3", "d4"):
        a = c.post("/accounts", json={"websites": ["x.com"], "username": user}, headers=h).json()
        c.post(f"/cappers/{teal['id']}/accounts", json={"account_id": a["id"]}, headers=h)
    rows = [
        ["Book", "Account", "Figure", "Status"],                              # no password column
        ["x.com", "a1", 10, "Active"],
        [],
        ["Book", "Account", "Password", "CHECK"],                             # no figure column: skipped
        ["x.com", "b2", "Flood", "ok"],
        [],
        ["Author", "Book", "Account", "Password", "Figure", "CHECK"],         # password before the figure
        ["me", "x.com", "c3", "Flood", -5370.0, None],
        ["me", "x.com", "d4", 123, 4.5, None],                                # a numeric password is not the figure
    ]
    r = c.post("/cappers/figures/import", data={"capper_id": teal["id"]},
               files={"file": ("Teal.xlsx", sheet_file({"2026-10-06 to 2026-10-12": rows}))}, headers=h)
    assert r.status_code == 200, r.text
    saved = {f["username"]: f["figure"] for f in c.get("/cappers/figures/all", headers=h).json()}
    assert saved == {"a1": "10.00", "c3": "-5370.00", "d4": "4.50"}           # b2's table has no figure column


def test_template_round_trip():
    from datetime import timedelta
    from openpyxl import load_workbook

    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    a = c.post("/accounts", json={"websites": ["trylnk.com"], "username": "lnk98213"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": a["id"]}, headers=h)

    tpl = c.get("/cappers/figures/import/template", headers=h)
    assert tpl.status_code == 200
    wb = load_workbook(io.BytesIO(tpl.content))
    assert len(wb.sheetnames) == 2
    ws = wb[wb.sheetnames[0]]
    assert [x.value for x in ws[1]] == ["Website", "Username", "Figure"]
    assert all(x.comment for x in ws[1])                       # each header explains itself
    ws.append(["trylnk.com", "lnk98213", -250.5])
    ws.append(["other.com", "ghost", 10])
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/cappers/figures/import", data={"capper_id": yellow["id"]},
               files={"file": ("Yellow.xlsx", buf.getvalue())}, headers=h)
    assert r.status_code == 200, r.text
    s = r.json()["sheets"][0]
    monday = date.today() - timedelta(days=date.today().weekday())
    assert (s["start_date"], s["end_date"]) == (str(monday), str(monday + timedelta(days=6)))
    assert (s["created"], s["unregistered_rows"]) == (1, 1)
    empty = r.json()["sheets"][1]                              # the empty second week saves nothing
    assert (empty["sheet_name"], empty["created"], empty["unregistered_rows"]) == (wb.sheetnames[1], 0, 0)
    assert len(c.get("/cappers/figures/all", headers=h).json()) == 1


def test_capper_figure_total(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    blue = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    partner = c.post("/partners", json={"name": "Baws"}, headers=h).json()
    ids = {}
    for u in ("a1", "a2"):
        acc = c.post("/accounts", json={"websites": ["x.com"], "username": u}, headers=h).json()
        ids[u] = acc["id"]
        for cap in (yellow, blue):
            c.post(f"/cappers/{cap['id']}/accounts", json={"account_id": acc["id"]}, headers=h)

    def upload(cap, name, wk, rows):
        return c.post("/cappers/figures/import", data={"capper_id": cap["id"]},
                      files={"file": (name, sheet_file({wk: [["Username", "Figure"], *rows]}))}, headers=h)

    # nothing imported yet: a total of zero, not an error
    d = c.get(f"/cappers/{yellow['id']}", headers=h).json()
    assert (d["figure_total"], d["figure_count"], d["figure_weeks"]) == ("0", 0, 0) or d["figure_total"] in ("0", "0.00")
    assert upload(yellow, "Yellow.xlsx", "10-6 - 10-12", [["a1", 100], ["a2", -30.5]]).status_code == 200
    assert upload(yellow, "Yellow.xlsx", "9-29 - 10-5", [["a1", 10]]).status_code == 200
    assert upload(blue, "Blue.xlsx", "10-6 - 10-12", [["a1", 1000]]).status_code == 200   # another capper: not counted
    # re-importing a week replaces the figure (one figure per account and week), it does not add
    assert upload(yellow, "Yellow.xlsx", "10-6 - 10-12", [["a1", 120], ["a2", -30.5]]).status_code == 200

    d = c.get(f"/cappers/{yellow['id']}", headers=h).json()
    assert d["figure_total"] == "99.50" and d["figure_count"] == 3 and d["figure_weeks"] == 2   # 120 - 30.5 + 10
    assert d["figure_start_date"] is None and d["figure_end_date"] is None
    w = c.get(f"/cappers/{yellow['id']}", params={"start_date": "2026-10-06", "end_date": "2026-10-12"}, headers=h).json()
    assert (w["figure_total"], w["figure_count"], w["figure_weeks"]) == ("89.50", 2, 1)
    assert (w["figure_start_date"], w["figure_end_date"]) == ("2026-10-06", "2026-10-12")
    other = c.get(f"/cappers/{blue['id']}", headers=h).json()
    assert other["figure_total"] == "1000.00"
    # a week without figures: zero; partners have no figure total
    z = c.get(f"/cappers/{yellow['id']}", params={"start_date": "2026-01-05", "end_date": "2026-01-11"}, headers=h).json()
    assert z["figure_count"] == 0 and float(z["figure_total"]) == 0
    p = c.get(f"/partners/{partner['id']}", headers=h).json()
    assert p["figure_total"] is None and p["figure_count"] == 0
    assert c.get(f"/cappers/{yellow['id']}", params={"start_date": "2026-10-06"}, headers=h).status_code == 422


def test_capper_total_with_unregistered(monkeypatch):
    import app.routers.imports as imports_module

    monkeypatch.setattr(imports_module, "parse_week", lambda name: parse_week(name, as_of=date(2026, 10, 13)))
    c, h = make_client()
    yellow = c.post("/cappers", json={"name": "Yellow"}, headers=h).json()
    blue = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    acc = c.post("/accounts", json={"websites": ["x.com"], "username": "a1"}, headers=h).json()
    c.post(f"/cappers/{yellow['id']}/accounts", json={"account_id": acc["id"]}, headers=h)
    c.post(f"/cappers/{blue['id']}/accounts", json={"account_id": acc["id"]}, headers=h)    # a1 is Blue's too
    c.post("/accounts", json={"websites": ["y.com"], "username": "bluesonly"}, headers=h)  # exists, not Yellow's

    def upload(cap, name, wk, rows):
        return c.post("/cappers/figures/import", data={"capper_id": cap["id"]},
                      files={"file": (name, sheet_file({wk: [["Username", "Figure"], *rows]}))}, headers=h)

    assert upload(yellow, "Yellow.xlsx", "10-6 - 10-12",
                  [["a1", 100], ["ghost1", 7], ["ghost2", -2.5], ["bluesonly", 10]]).status_code == 200
    assert upload(yellow, "Yellow.xlsx", "9-29 - 10-5", [["a1", 5], ["ghost1", 1]]).status_code == 200
    assert upload(blue, "Blue.xlsx", "10-6 - 10-12", [["a1", 3], ["ghost1", 500]]).status_code == 200  # not Yellow's

    d = c.get(f"/cappers/{yellow['id']}", headers=h).json()
    assert (d["figure_total"], d["figure_count"]) == ("105.00", 2)
    assert (d["unregistered_figure_total"], d["unregistered_figure_count"]) == ("15.50", 4)   # 7 - 2.5 + 10 + 1
    assert d["figure_total_all"] == "120.50"
    w = c.get(f"/cappers/{yellow['id']}", params={"start_date": "2026-10-06", "end_date": "2026-10-12"}, headers=h).json()
    assert (w["figure_total"], w["unregistered_figure_total"], w["unregistered_figure_count"], w["figure_total_all"]) == (
        "100.00", "14.50", 3, "114.50")
    z = c.get(f"/cappers/{yellow['id']}", params={"start_date": "2026-01-05", "end_date": "2026-01-11"}, headers=h).json()
    assert float(z["figure_total_all"]) == 0 and z["unregistered_figure_count"] == 0
    # a re-import of the week replaces that week's unregistered rows instead of adding to them
    assert upload(yellow, "Yellow.xlsx", "10-6 - 10-12", [["a1", 100], ["ghost1", 7]]).status_code == 200
    assert c.get(f"/cappers/{yellow['id']}", headers=h).json()["unregistered_figure_total"] == "8.00"
    assert c.get(f"/partners/{c.post('/partners', json={'name': 'P'}, headers=h).json()['id']}", headers=h).json()["figure_total_all"] is None
