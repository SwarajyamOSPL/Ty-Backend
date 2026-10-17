import os

os.environ["DATABASE_URL"] = "sqlite:///./test.db"

from fastapi.testclient import TestClient  # noqa: E402

from app.db import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.security import seed_admin  # noqa: E402


def test_full_flow():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_admin(db)
        seed_admin(db)  # idempotent: still exactly one row
    c = TestClient(app)
    r = c.post("/auth/login", data={"username": "admin", "password": "nope"})
    assert r.status_code == 401 and r.json()["detail"] == "Invalid user name and/or password"
    assert c.post("/auth/login", data={"username": "who", "password": "x"}).json()["detail"] == "Invalid user name and/or password"
    assert c.get("/partners").status_code == 401
    tok = c.post("/auth/login", data={"username": "admin", "password": "admin"}).json()
    assert tok["username"] == "admin"
    h = {"Authorization": f"Bearer {tok['access_token']}"}

    # partner names may repeat (two different partners with the same name); each has its own id
    p1 = c.post("/partners", json={"name": "Alpha"}, headers=h).json()
    zz1 = c.post("/partners", json={"name": "Zz  Top"}, headers=h)
    zz2 = c.post("/partners", json={"name": " zz   TOP "}, headers=h)
    assert zz1.status_code == 201 and zz2.status_code == 201 and zz1.json()["id"] != zz2.json()["id"]
    assert (zz1.json()["name"], zz2.json()["name"]) == ("Zz Top", "zz TOP")  # spaces trimmed/collapsed
    assert [x["name"] for x in c.get("/partners", params={"prefix": "zz"}, headers=h).json()] == ["Zz Top", "zz TOP"]
    p2 = c.post("/partners", json={"name": "Bravo"}, headers=h).json()
    for i in range(105):  # paged: 50 by default, X-* headers say how many there are
        c.post("/partners", json={"name": f"Bulk {i:03d}"}, headers=h)
    r = c.get("/partners", headers=h)
    assert len(r.json()) == 50 and r.headers["X-Total-Count"] == str(4 + 105)
    assert (r.headers["X-Page"], r.headers["X-Page-Size"], r.headers["X-Total-Pages"]) == ("1", "50", "3")
    seen = []
    for page in (1, 2, 3, 4):  # every partner exactly once across the pages; past the end is empty
        seen += [x["id"] for x in c.get("/partners", params={"page": page, "page_size": 50}, headers=h).json()]
    assert len(seen) == len(set(seen)) == 4 + 105
    assert len(c.get("/partners", params={"page_size": 500}, headers=h).json()) == 4 + 105
    assert len(c.get("/partners", params={"limit": 5}, headers=h).json()) == 5  # old name of page_size
    assert c.get("/partners", params={"limit": 0}, headers=h).status_code == 422
    assert c.get("/partners", params={"page": 0}, headers=h).status_code == 422
    assert c.get("/partners", params={"page_size": 501}, headers=h).status_code == 422
    assert [x["name"] for x in c.get("/partners", params={"search": "ulk 10"}, headers=h).json()] == \
        [f"Bulk {i}" for i in range(100, 105)]  # "contains", any case
    # capper names stay unique
    assert c.post("/cappers", json={"name": "  Dark   Red "}, headers=h).json()["name"] == "Dark Red"  # new colors are fine
    cap = c.post("/cappers", json={"name": " lime  GREEN "}, headers=h)
    assert cap.status_code == 201 and cap.json()["name"] == "Lime Green"
    assert c.post("/cappers", json={"name": "Lime Green"}, headers=h).status_code == 409
    colors = c.get("/cappers/colors", headers=h).json()
    assert len(colors) == 16 and [x["name"] for x in colors if x["created"]] == ["Lime Green"]
    assert [p["name"] for p in c.get("/partners?prefix=a", headers=h).json()] == ["Alpha"]

    # accounts: password never listed
    a1 = c.post("/accounts", json={"websites": ["site1.com"], "username": "u1", "password": "pw1"},
                headers=h).json()
    a2 = c.post("/accounts", json={"websites": ["site2.com"]}, headers=h).json()
    full = c.post("/accounts", json={"websites": ["s3.com"], "balance": "-12.5", "max_win": "500", "city": " Pune "}, headers=h).json()
    assert (full["balance"], full["max_win"], full["city"]) == ("-12.50", "500.00", "Pune")
    assert a2["balance"] is None and a2["city"] is None
    assert c.post("/accounts", json={"websites": ["x"], "max_win": "-1"}, headers=h).status_code == 422
    # balance and max win are numbers, or text kept as written ("Infinite", "Unlimited", "No max")
    ex = c.post("/accounts", json={"websites": ["Playbetnow.com"], "username": "Db6211", "password": "bang",
                                   "balance": "30000", "max_win": "Unlimited", "city": "Dallas"}, headers=h)
    assert ex.status_code == 201 and ex.json()["max_win"] == "Unlimited" and ex.json()["balance"] == "30000.00", ex.text
    assert c.get(f"/accounts/{ex.json()['id']}", headers=h).json()["max_win"] == "Unlimited"
    assert any(x["max_win"] == "Unlimited" for x in c.get("/accounts", headers=h).json())
    for raw, want in (("  unlimited ", "Unlimited"), ("UNLIMITED", "Unlimited"), ("1,500", "1500.00"), (750, "750.00"),
                      (None, None), ("", None), ("No max", "No max"), ("  no    max ", "no max"), ("abc", "abc"),
                      ("NaN", "NaN"), ("Infinity", "Infinity")):
        r = c.post("/accounts", json={"websites": ["mw.com"], "max_win": raw}, headers=h)
        assert r.status_code == 201 and r.json()["max_win"] == want, (raw, r.text)
    for raw in ("-5", -5, True, "x" * 101, float("nan") if False else "1" + "x" * 100):  # negative number, boolean, too long
        assert c.post("/accounts", json={"websites": ["mw.com"], "max_win": raw}, headers=h).status_code == 422, raw
    for raw, want in (("Infinite ", "Infinite"), ("-12.5", "-12.50"), ("50,000", "50000.00"), (0, "0.00"), ("", None)):
        r = c.post("/accounts", json={"websites": ["bal.com"], "balance": raw}, headers=h)
        assert r.status_code == 201 and r.json()["balance"] == want, (raw, r.text)
    assert c.post("/accounts", json={"websites": ["bal.com"], "balance": True}, headers=h).status_code == 422
    # the text shows exactly as written everywhere and can sit next to a numeric max win
    t = c.post("/accounts", json={"websites": ["txt.com"], "username": "tt", "password": "p", "balance": "Infinite", "max_win": "No max"}, headers=h).json()
    assert (t["balance"], t["max_win"]) == ("Infinite", "No max")
    assert c.get(f"/accounts/{t['id']}", headers=h).json()["balance"] == "Infinite"
    # same credentials again: the shared account keeps its values (text counts as a value)
    again = c.post("/accounts", json={"websites": ["txt.com"], "username": "tt", "password": "p", "balance": "5", "max_win": "9"}, headers=h)
    assert again.status_code == 200 and (again.json()["balance"], again.json()["max_win"]) == ("Infinite", "No max")
    multi = c.post("/accounts", json={"websites": [" a.com ", "b.com", "A.com"], "username": "mu",
                                      "password": "mp", "city": "Delhi"}, headers=h)
    assert multi.status_code == 201 and multi.json()["websites"] == ["a.com", "b.com"]  # one account
    assert multi.json()["username"] == "mu" and multi.json()["city"] == "Delhi"
    assert [a["id"] for a in c.get("/accounts?website=B.", headers=h).json()] == [multi.json()["id"]]
    assert c.post("/accounts", json={"websites": []}, headers=h).status_code == 422
    assert c.post("/accounts", json={"websites": ["  "]}, headers=h).status_code == 422
    assert c.post("/accounts", json={"website": "old.com"}, headers=h).status_code == 422
    one = c.get(f"/accounts/{multi.json()['id']}", headers=h)
    assert one.status_code == 200 and one.json() == multi.json() and "password" not in one.json()
    assert c.get("/accounts/99999", headers=h).status_code == 404
    assert c.get("/accounts/abc", headers=h).status_code == 422
    assert c.get(f"/accounts/{multi.json()['id']}").status_code == 401
    listed = c.get("/accounts", headers=h).json()
    assert all("password" not in a for a in listed)

    # assign: no duplicate per partner, but same account under another partner is fine
    assert c.post(f"/partners/{p1['id']}/accounts", json={"account_id": a1["id"]},
                  headers=h).status_code == 201
    assert c.post(f"/partners/{p1['id']}/accounts", json={"account_id": a1["id"]},
                  headers=h).status_code == 409
    assert c.post(f"/partners/{p2['id']}/accounts", json={"account_id": a1["id"]},
                  headers=h).status_code == 201
    c.post(f"/partners/{p1['id']}/accounts", json={"account_id": a2["id"]}, headers=h)
    assert len(c.get(f"/partners/{p1['id']}/accounts", headers=h).json()) == 2
    assert len(c.get(f"/partners/{p2['id']}/accounts?prefix=site2", headers=h).json()) == 0

    # enter figure: credentials pre-fill, missing ones come back null
    cred = c.get(f"/partners/{p1['id']}/accounts/{a1['id']}/credentials", headers=h).json()
    assert cred["username"] == "u1" and cred["password"] == "pw1"
    cred2 = c.get(f"/partners/{p1['id']}/accounts/{a2['id']}/credentials", headers=h).json()
    assert cred2["username"] is None and cred2["password"] is None

    r = c.post("/partners/figures",
               json={"partner_id": p1["id"], "account_id": a2["id"], "figure": "150.50",
                     "username": "new", "password": "np"},
               headers=h)
    assert r.status_code == 201 and r.json()["username"] == "new"
    cred2 = c.get(f"/partners/{p1['id']}/accounts/{a2['id']}/credentials", headers=h).json()
    assert cred2["password"] == "np"
    # figure on an unassigned account is rejected
    assert c.post("/partners/figures", json={"partner_id": p2["id"], "account_id": a2["id"], "figure": "1"},
                  headers=h).status_code == 404

    assert c.get("/partners/figures/all", headers=h).json()[0]["figure"] == "150.50"
    cf = c.post("/cappers/figures", json={"capper_id": cap.json()["id"], "account_id": a1["id"],
                                         "figure": "5", "start_date": "2026-10-06",
                                         "end_date": "2026-10-12"}, headers=h)
    assert cf.status_code == 404  # account not assigned to that capper
    c.post(f"/cappers/{cap.json()['id']}/accounts", json={"account_id": a1["id"]}, headers=h)
    body = {"capper_id": cap.json()["id"], "account_id": a1["id"], "figure": "5"}
    assert c.post("/cappers/figures", json=body, headers=h).status_code == 422  # dates required
    bad = {**body, "start_date": "2026-10-12", "end_date": "2026-10-06"}
    assert c.post("/cappers/figures", json=bad, headers=h).status_code == 422
    ok = c.post("/cappers/figures", json={**body, "start_date": "2026-10-06",
                                          "end_date": "2026-10-12"}, headers=h)
    assert ok.status_code == 201 and ok.json()["start_date"] == "2026-10-06"

    d = c.get("/dashboard", headers=h).json()
    assert d["partners"] == 109 and d["cappers"] == 2 and d["partner_figure_total"] == "150.50"


def test_excel_import(monkeypatch):
    import io
    from datetime import date as _date

    import app.routers.imports as imports_module
    from app.week_names import parse_week as real_parse_week

    # sheet names have no year: pin "today" so the year they resolve to is deterministic
    monkeypatch.setattr(imports_module, "parse_week", lambda name: real_parse_week(name, as_of=_date(2026, 10, 13)))

    from openpyxl import Workbook, load_workbook

    from app.week_names import parse_week

    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_admin(db)
    c = TestClient(app)
    tok = c.post("/auth/login", data={"username": "admin", "password": "admin"}).json()
    h = {"Authorization": f"Bearer {tok['access_token']}"}
    cap = c.post("/cappers", json={"name": "Red"}, headers=h).json()
    other = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    a1 = c.post("/accounts", json={"websites": ["Site1.com"], "username": "u1"}, headers=h).json()
    a2 = c.post("/accounts", json={"websites": ["site2.com"], "username": "u2"}, headers=h).json()
    a3 = c.post("/accounts", json={"websites": ["m1.com", "m2.com"], "username": "mu"}, headers=h).json()
    nouser = c.post("/accounts", json={"websites": ["nu.com"]}, headers=h).json()
    for a in (a1, a2, a3, nouser):
        c.post(f"/cappers/{cap['id']}/accounts", json={"account_id": a["id"]}, headers=h)

    HEAD = ("username", "figure")

    def book(sheets):
        """sheets: {sheet name: rows} or {sheet name: (header, rows)}"""
        wb = Workbook()
        wb.remove(wb.active)
        for name, spec in sheets.items():
            header, rows = spec if isinstance(spec, tuple) else (HEAD, spec)
            ws = wb.create_sheet(name)
            ws.append(list(header))
            for r in rows:
                ws.append(list(r))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def upload(content, capper_id=cap["id"], name="Red.xlsx"):
        data = {"capper_id": capper_id}
        return c.post("/cappers/figures/import", headers=h, data=data, files={"file": (name, content)})

    def errors(r):
        return r.json()["detail"]["errors"]

    def figs():
        return c.get("/cappers/figures/all", headers=h).json()

    tpl = c.get("/cappers/figures/import/template", headers=h)
    assert tpl.status_code == 200
    assert parse_week(load_workbook(io.BytesIO(tpl.content)).sheetnames[0])  # template sheet name is a readable week

    # several weeks in one file; the sheet name is the week; a notes sheet is ignored
    week1 = [("U1", 100), ("u2", "-25.5"), (" mu ", 7), ("someone_else", 1), ("another", 2), (None, 3)]
    week2 = [("u1", 11), ("u2", 12)]
    good = book({"10-6 - 10-12": week1, "Sept 29 - Oct 5": week2, "Notes": (("a", "b"), [("x", "y")])})
    r = upload(good)
    assert r.status_code == 200, r.text
    j = r.json()
    assert (j["created"], j["updated"], j["unregistered_rows"], j["ignored_rows"]) == (5, 0, 2, 1) and j["ignored_sheets"] == ["Notes"]
    by_sheet = {x["sheet_name"]: x for x in j["sheets"]}
    w1 = by_sheet["10-6 - 10-12"]
    assert (w1["start_date"], w1["end_date"]) == ("2026-10-06", "2026-10-12")
    assert (w1["created"], w1["unregistered_rows"], w1["ignored_rows"]) == (3, 2, 1)
    w2 = by_sheet["Sept 29 - Oct 5"]
    assert (w2["start_date"], w2["end_date"]) == ("2026-09-29", "2026-10-05")
    saved = figs()
    assert len(saved) == 5 and {(f["start_date"], f["end_date"]) for f in saved} == {
        ("2026-10-06", "2026-10-12"), ("2026-09-29", "2026-10-05")}
    assert sorted(f["figure"] for f in saved if f["start_date"] == "2026-10-06") == ["-25.50", "100.00", "7.00"]

    # the same file again overwrites each week; names in other formats work too
    r = upload(good, name="red_weeks_41_42.xlsx")
    assert (r.json()["created"], r.json()["updated"]) == (0, 5) and len(figs()) == 5
    r = upload(book({"106 1011": [("u1", 5)], "2026-10-13 to 2026-10-19": [("u2", 6)]}))
    assert r.status_code == 200
    assert [(x["start_date"], x["end_date"]) for x in r.json()["sheets"]] == [
        ("2026-10-06", "2026-10-11"), ("2026-10-13", "2026-10-19")]
    # no year in the sheet name: it comes from today (pinned to 2026-10-13 here), never from the future
    assert upload(book({"10-6 - 10-12": [("u1", 1)]})).json()["sheets"][0]["start_date"] == "2026-10-06"
    assert upload(book({"11-24 - 11-30": [("u1", 1)]})).json()["sheets"][0]["start_date"] == "2025-11-24"
    assert upload(book({"2024-10-06 to 2024-10-12": [("u1", 1)]})).json()["sheets"][0]["start_date"] == "2024-10-06"  # explicit year wins

    # file name must name the selected capper
    r = upload(good, name="Blue.xlsx")
    assert r.status_code == 422 and "'Blue'" in errors(r)[0] and "'Red'" in errors(r)[0]
    r = upload(good, name="Turquoise.xlsx", capper_id=other["id"])  # a color that is not the selected capper (Blue)
    assert r.status_code == 422 and "'Turquoise'" in errors(r)[0] and "'Blue'" in errors(r)[0]
    c.post("/cappers", json={"name": "Light Blue"}, headers=h)  # "Blue" is also inside "Light Blue": the longer name wins
    assert upload(good, name="Light Blue.xlsx", capper_id=other["id"]).status_code == 422
    assert upload(good, name="f.xlsx").status_code == 422  # no capper name in file name
    assert upload(good, name="Reddish.xlsx").status_code == 422  # whole words only
    assert upload(good, name="Red - week 41 (final).xlsx").status_code == 200
    assert upload(good, capper_id=other["id"], name="Blue.xlsx").status_code == 422  # accounts not assigned
    assert upload(good, name="Red.csv").status_code == 422
    assert upload(b"not excel").status_code == 422

    # sheet-name problems reject the whole file
    before = len(figs())
    r = upload(book({"10-6 - 10-12": [("u1", 77)], "Week 41": [("u1", 1)]}))
    assert r.status_code == 422 and "Week 41" in errors(r)[0] and len(figs()) == before
    assert not any(f["figure"] == "77.00" for f in figs())  # the good sheet was not saved either
    r = upload(book({"10-6 - 10-12": [("u1", 1)], "Oct 6 - Oct 12": [("u2", 2)]}))
    assert r.status_code == 422 and "same week" in errors(r)[0]
    r = upload(book({"Notes": (("a", "b"), [(1, 2)])}))
    assert r.status_code == 422 and "username" in errors(r)[0]
    assert upload(book({"10-6 - 10-12": (("website", "figure"), [(1, 2)])})).status_code == 422  # username required

    # all-or-nothing across sheets: a bad figure anywhere rejects everything
    bad = book({"10-6 - 10-12": [("u1", 5), ("u2", "abc"), ("u1", 4)], "9-29 - 10-5": [("u1", 8)]})
    r = upload(bad)
    assert r.status_code == 422 and len(errors(r)) == 1 and "[10-6 - 10-12]" in errors(r)[0]
    assert len(figs()) == before
    # the same username again with the SAME figure is ignored; with a different figure the figures are added
    r = upload(book({"10-6 - 10-12": [("u1", 5), ("u1", 5)]}))
    assert r.status_code == 200 and r.json()["sheets"][0]["repeated_ignored"] == 1 and r.json()["sheets"][0]["created"] + r.json()["sheets"][0]["updated"] == 1
    r = upload(book({"10-6 - 10-12": [("u1", 5), ("u1", 5.0), ("u1", "5.00")]}))
    assert r.status_code == 200 and r.json()["sheets"][0]["repeated_ignored"] == 2
    r = upload(book({"10-6 - 10-12": [("u1", 5), ("u1", 6), ("u1", 5), ("u1", -2)]}))
    sh = r.json()["sheets"][0]
    assert r.status_code == 200 and (sh["summed_rows"], sh["repeated_ignored"]) == (2, 1)
    assert [f["figure"] for f in figs() if f["username"] == "u1" and f["start_date"] == sh["start_date"]].count("9.00") == 1  # 5 + 6 - 2; the second 5 is a repeat
    # nothing matches in any sheet -> rejected (probably the wrong file)
    r = upload(book({"10-6 - 10-12": [("x1", 1), ("x2", 2)]}))
    assert r.status_code == 422 and "2 row(s) not saved" in errors(r)[0]
    # a sheet with no matching rows is fine as long as another sheet matches
    r = upload(book({"10-6 - 10-12": [("x1", 1)], "9-29 - 10-5": [("u1", 9)]}))
    assert r.status_code == 200 and sorted(x["created"] + x["updated"] for x in r.json()["sheets"]) == [0, 1]


def test_site_figure():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_admin(db)
    c = TestClient(app)
    tok = c.post("/auth/login", data={"username": "admin", "password": "admin"}).json()
    h = {"Authorization": f"Bearer {tok['access_token']}"}
    p1 = c.post("/partners", json={"name": "P1"}, headers=h).json()
    p2 = c.post("/partners", json={"name": "P2"}, headers=h).json()
    red = c.post("/cappers", json={"name": "Red"}, headers=h).json()
    blue = c.post("/cappers", json={"name": "Blue"}, headers=h).json()
    green = c.post("/cappers", json={"name": "Green"}, headers=h).json()
    acct = c.post("/accounts", json={"websites": ["s.com"], "username": "Uname"}, headers=h).json()
    solo = c.post("/accounts", json={"websites": ["t.com"], "username": "lonely"}, headers=h).json()
    for party in (p1, p2):
        c.post(f"/partners/{party['id']}/accounts", json={"account_id": acct["id"]}, headers=h)
    c.post(f"/partners/{p1['id']}/accounts", json={"account_id": solo["id"]}, headers=h)
    for capper in (red, blue, green):
        c.post(f"/cappers/{capper['id']}/accounts", json={"account_id": acct["id"]}, headers=h)
    wk = {"start_date": "2026-10-06", "end_date": "2026-10-12"}
    base = f"/partners/{p1['id']}/accounts/{acct['id']}"

    # username search among the partner's accounts
    assert [a["id"] for a in c.get(f"/partners/{p1['id']}/accounts?username=un", headers=h).json()] == [acct["id"]]
    assert c.get(f"/partners/{p1['id']}/accounts?username=zz", headers=h).json() == []
    # cappers assigned to the account
    assert [x["capper_name"] for x in c.get(f"{base}/cappers", headers=h).json()] == ["Blue", "Green", "Red"]
    assert c.get(f"/partners/{p1['id']}/accounts/{solo['id']}/cappers", headers=h).json() == []

    # nothing imported yet -> cannot save
    assert c.post(f"{base}/site-figure", json=wk, headers=h).status_code == 422
    # Red and Blue have a figure for the week, Green does not (different week)
    for capper, fig, week in ((red, "100.25", wk), (blue, "-40", wk),
                              (green, "999", {"start_date": "2026-09-29", "end_date": "2026-10-05"})):
        r = c.post("/cappers/figures", json={"capper_id": capper["id"], "account_id": acct["id"],
                                             "figure": fig, **week}, headers=h)
        assert r.status_code == 201
    # a second, corrected entry for Red in the same week: the latest wins
    c.post("/cappers/figures", json={"capper_id": red["id"], "account_id": acct["id"],
                                     "figure": "110.25", **wk}, headers=h)

    v = c.get(f"{base}/site-figure", params=wk, headers=h).json()
    assert v["total"] == "70.25" and v["missing_cappers"] == ["Green"] and v["saved"] is None
    assert {x["capper_name"]: x["figure"] for x in v["cappers"]} == {"Blue": "-40.00", "Green": None, "Red": "110.25"}

    v = c.post(f"{base}/site-figure", json=wk, headers=h).json()
    assert v["saved"]["site_figure"] == "70.25"
    first_id = v["saved"]["id"]
    # Green's sheet arrives later -> recalculating overwrites the same row
    c.post("/cappers/figures", json={"capper_id": green["id"], "account_id": acct["id"],
                                     "figure": "30", **wk}, headers=h)
    v = c.post(f"{base}/site-figure", json=wk, headers=h).json()
    assert v["total"] == "100.25" and v["missing_cappers"] == [] and v["saved"]["id"] == first_id
    saved = c.get("/site-figures", headers=h).json()
    assert len(saved) == 1 and saved[0]["site_figure"] == "100.25" and saved[0]["username"] == "Uname"
    # one Site Figure per ACCOUNT, shared by every partner that has the account
    assert sorted(x["name"] for x in saved[0]["partners"]) == ["P1", "P2"]
    assert len(c.get("/site-figures", params={"partner_id": p2["id"]}, headers=h).json()) == 1
    assert c.get("/site-figures", params={"partner_id": 99999}, headers=h).json() == []
    r = c.post(f"/partners/{p2['id']}/accounts/{acct['id']}/site-figure", json=wk, headers=h)
    assert r.status_code == 200 and r.json()["saved"]["id"] == first_id  # same row, not a second one
    assert len(c.get("/site-figures", headers=h).json()) == 1

    # errors: account with no capper, account not assigned to partner, bad range
    assert c.post(f"/partners/{p1['id']}/accounts/{solo['id']}/site-figure", json=wk, headers=h).status_code == 422
    assert c.get(f"/partners/{p2['id']}/accounts/{solo['id']}/site-figure", params=wk, headers=h).status_code == 404
    bad = {"start_date": "2026-10-12", "end_date": "2026-10-06"}
    assert c.post(f"{base}/site-figure", json=bad, headers=h).status_code == 422
    assert c.get(f"{base}/site-figure", params=bad, headers=h).status_code == 422
    assert c.get(f"{base}/site-figure", params=wk).status_code == 401
